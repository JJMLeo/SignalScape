# -*- coding: utf-8 -*-
"""生成 tests/fixtures/ 下的测试样本文件。

用法（推荐在 SignalScapePy38 虚拟环境中运行）::

    python tests/make_fixtures.py            # 生成 tdms（需 nptdms）+ xlsx（需 openpyxl）
    python tests/make_fixtures.py --xlsx-only  # 只生成 xlsx

虚拟环境未安装 openpyxl 时，可用基础解释器单独生成 xlsx::

    <base python> tests/make_fixtures.py --xlsx-only
"""
import os
import sys

FIXTURE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')
N_POINTS = 1000


def _ensure_dir():
    if not os.path.isdir(FIXTURE_DIR):
        os.makedirs(FIXTURE_DIR)


def make_tdms():
    """生成 sample.tdms：1 个组 + 4 个通道（其中 2 个应被默认过滤规则排除）"""
    from nptdms import TdmsWriter, RootObject, GroupObject, ChannelObject

    path = os.path.join(FIXTURE_DIR, 'sample.tdms')
    if os.path.exists(path):
        os.remove(path)

    t = [i * 0.02 for i in range(N_POINTS)]
    import math
    temp = [50.0 + 10.0 * math.sin(i / 50.0) for i in range(N_POINTS)]
    press = [100.0 + 5.0 * math.cos(i / 30.0) for i in range(N_POINTS)]
    status = [0] * N_POINTS
    fault = [1.0] * N_POINTS

    with TdmsWriter(path) as writer:
        writer.write_segment([
            RootObject(properties={'name': 'sample'}),
            GroupObject('试验组'),
            ChannelObject('试验组', '温度', temp),
            ChannelObject('试验组', '压力', press),
            ChannelObject('试验组', '转速_status', status),
            ChannelObject('试验组', '故障标志', fault),
        ])
    print('written', path)
    return path


def make_xlsx():
    """生成 sample.xlsx（常规时刻）与 midnight.xlsx（跨零点时刻）"""
    try:
        from openpyxl import Workbook
    except ImportError:
        print('skip xlsx: openpyxl not available in this interpreter')
        return None

    # 约定：第1行日期；第3行变量编号；第4行变量名称；第5行单位；第6行起数据
    def build(path, times, values, names):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value='2024-05-01')
        for c, nm in enumerate(names, start=2):
            ws.cell(row=3, column=c, value='P%d' % c)
            ws.cell(row=4, column=c, value=nm)
            ws.cell(row=5, column=c, value='-')
        for r, (t, vals) in enumerate(zip(times, values), start=6):
            ws.cell(row=r, column=1, value=t)
            for c, v in enumerate(vals, start=2):
                if v is None:
                    continue  # 空单元格 → 测试 NaN 处理
                ws.cell(row=r, column=c, value=v)
        wb.save(path)
        print('written', path)

    n = 50
    times = ['%02d:%02d:%02d' % (0, 0, i) for i in range(n)]
    values = []
    for i in range(n):
        row = [10.0 + i, 20.0 + 2 * i, 1]
        if i == 10:  # 制造一个空单元格
            row[1] = None
        values.append(row)
    build(os.path.join(FIXTURE_DIR, 'sample.xlsx'), times, values,
          ['温度', '压力', '状态_status'])

    # 跨零点：23:59:58 → 00:00:03
    times2 = ['23:59:58', '23:59:59', '00:00:00', '00:00:01', '00:00:02', '00:00:03']
    values2 = [[float(i), 100.0 + i] for i in range(len(times2))]
    build(os.path.join(FIXTURE_DIR, 'midnight.xlsx'), times2, values2,
          ['温度', '压力'])
    return os.path.join(FIXTURE_DIR, 'sample.xlsx')


if __name__ == '__main__':
    _ensure_dir()
    if '--xlsx-only' not in sys.argv:
        make_tdms()
    make_xlsx()
